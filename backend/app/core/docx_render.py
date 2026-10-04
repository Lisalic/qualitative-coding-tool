"""Word (``.docx``) document builder for exports.

The default for exports a researcher reads, prints or pastes into a
write-up -- a codebook, a code report, memos, a saved summary. Every
piece of text goes in through the helpers here, which run it through
``export_text.xml_safe`` first: python-docx raises on the control
characters XML forbids, so a vertical tab pasted from Word would
otherwise turn an export into a 500.

``markdown_to_docx`` converts the markdown blobs this app stores (LLM
summaries and comparisons) at the block level the generators actually
emit. It never interprets raw HTML (it stays literal text) and never
creates a live hyperlink field -- a link becomes ``text (url)`` -- so
nothing in a stored blob turns into active content in Word.

Output is byte-deterministic (see ``xlsx_render.freeze_ooxml``).
"""

from __future__ import annotations

import io
import re
from typing import Any, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from backend.app.core.export_text import xml_safe
from backend.app.core.xlsx_render import CREATOR, FIXED_TIMESTAMP, freeze_ooxml

_MONOSPACE = "Courier New"


def _add_run(paragraph: Paragraph, text: Any, *, bold: bool = False, italic: bool = False, mono: bool = False) -> None:
    run = paragraph.add_run(xml_safe(text))
    run.bold = bold or None
    run.italic = italic or None
    if mono:
        run.font.name = _MONOSPACE


def new_document(title: str, subtitle_lines: Sequence[str] = ()) -> DocumentType:
    """A document opened by a title and optional italic subtitle lines."""
    doc = Document()
    doc.core_properties.title = xml_safe(title)
    doc.add_heading(xml_safe(title), level=0)
    for line in subtitle_lines:
        _add_run(doc.add_paragraph(), line, italic=True)
    return doc


def add_heading(doc: DocumentType, text: Any, level: int) -> Paragraph:
    return doc.add_heading(xml_safe(text), level=max(1, min(level, 9)))


def add_paragraph(
    doc: DocumentType,
    text: Any,
    *,
    label: str | None = None,
    style: str | None = None,
    italic: bool = False,
) -> Paragraph:
    """One paragraph, optionally led by a bold ``label: ``."""
    paragraph = doc.add_paragraph(style=style)
    if label:
        _add_run(paragraph, f"{label}: ", bold=True)
    _add_run(paragraph, text, italic=italic)
    return paragraph


def add_table(doc: DocumentType, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    """A bordered table with a bold header row."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        _add_run(cell.paragraphs[0], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            _add_run(cell.paragraphs[0], "" if value is None else value)
    doc.add_paragraph()


def to_bytes(doc: DocumentType) -> bytes:
    props = doc.core_properties
    props.author = CREATOR
    props.last_modified_by = CREATOR
    props.comments = ""
    props.revision = 1
    props.created = FIXED_TIMESTAMP
    props.modified = FIXED_TIMESTAMP
    props.last_printed = FIXED_TIMESTAMP
    buffer = io.BytesIO()
    doc.save(buffer)
    return freeze_ooxml(buffer.getvalue())


# --- markdown -------------------------------------------------------------

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")
_RULE_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
_TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")
_INLINE_RE = re.compile(
    r"(`[^`]+`"                              # code
    r"|\*\*(?=\S).+?(?<=\S)\*\*"             # **bold**
    r"|__(?=\S).+?(?<=\S)__"                 # __bold__
    r"|\*(?=[^\s*]).+?(?<=[^\s*])\*"         # *italic*
    r"|(?<!\w)_(?=[^\s_]).+?(?<=[^\s_])_(?!\w)"  # _italic_
    r"|\[[^\]]+\]\([^)\s]+\))"               # [text](url)
)
_LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)\s]+)\)$")
_ESCAPE_RE = re.compile(r"\\([\\`*_{}\[\]()#+\-.!|<>~])")
# A backslash-escaped character is parked in the Private Use Area while
# inline markup is matched, so `\*` can never open or close emphasis.
_ESCAPE_BASE = 0xE000
_PARKED_RE = re.compile("[\ue000-\ue07f]")


def _park_escapes(text: str) -> str:
    return _ESCAPE_RE.sub(lambda m: chr(_ESCAPE_BASE + ord(m.group(1))), text)


def _unpark(text: str) -> str:
    return _PARKED_RE.sub(lambda m: chr(ord(m.group(0)) - _ESCAPE_BASE), text)


def _inline(paragraph: Paragraph, text: str, *, bold: bool = False) -> None:
    _add_inline(paragraph, _park_escapes(text), bold=bold)


def _add_inline(paragraph: Paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    for token in _INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`") and len(token) > 1:
            _add_run(paragraph, _unpark(token[1:-1]), bold=bold, italic=italic, mono=True)
        elif token[:2] in ("**", "__") and token[-2:] == token[:2] and len(token) > 4:
            _add_inline(paragraph, token[2:-2], bold=True, italic=italic)
        elif token[0] in "*_" and token[-1] == token[0] and len(token) > 2:
            _add_inline(paragraph, token[1:-1], bold=bold, italic=True)
        elif link := _LINK_RE.match(token):
            # Written as text: no hyperlink relationship is ever created
            # from stored content.
            _add_inline(paragraph, link.group(1), bold=bold, italic=italic)
            _add_run(paragraph, f" ({_unpark(link.group(2))})", bold=bold, italic=italic)
        else:
            _add_run(paragraph, _unpark(token), bold=bold, italic=italic)


def _split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|") and not stripped.endswith("\\|"):
        stripped = stripped[:-1]
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", stripped)]


def _add_markdown_table(doc: DocumentType, header: list[str], body: list[list[str]]) -> None:
    cols = len(header)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        _inline(cell.paragraphs[0], text, bold=True)
    for row in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, (row + [""] * cols)[:cols], strict=True):
            _inline(cell.paragraphs[0], text)
    doc.add_paragraph()


def markdown_to_docx(doc: DocumentType, markdown: str) -> None:
    """Append ``markdown`` to ``doc``: headings, paragraphs, bullet and
    numbered lists, block quotes, fenced code, pipe tables, rules, and
    inline bold/italic/code/links."""
    lines = (markdown or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            _inline(doc.add_paragraph(), " ".join(line.strip() for line in paragraph_lines))
            paragraph_lines.clear()

    while i < len(lines):
        line = lines[i]

        if _FENCE_RE.match(line):
            flush_paragraph()
            fence = _FENCE_RE.match(line).group(1)
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence (or end of input)
            paragraph = doc.add_paragraph()
            _add_run(paragraph, "\n".join(code), mono=True)
            paragraph.paragraph_format.left_indent = Pt(18)
            continue

        if not line.strip():
            flush_paragraph()
            i += 1
            continue

        if heading := _HEADING_RE.match(line):
            flush_paragraph()
            _inline(doc.add_heading("", level=len(heading.group(1))), heading.group(2))
            i += 1
            continue

        if _RULE_RE.match(line):
            flush_paragraph()
            doc.add_paragraph()
            i += 1
            continue

        if "|" in line and i + 1 < len(lines) and _TABLE_SEPARATOR_RE.match(lines[i + 1]) and "-" in lines[i + 1]:
            flush_paragraph()
            header = _split_table_row(line)
            i += 2
            body: list[list[str]] = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                body.append(_split_table_row(lines[i]))
                i += 1
            _add_markdown_table(doc, header, body)
            continue

        if _QUOTE_RE.match(line):
            flush_paragraph()
            quoted: list[str] = []
            while i < len(lines) and (match := _QUOTE_RE.match(lines[i])):
                quoted.append(match.group(1).strip())
                i += 1
            _inline(doc.add_paragraph(style="Quote"), " ".join(q for q in quoted if q))
            continue

        if bullet := _BULLET_RE.match(line):
            flush_paragraph()
            style = "List Bullet 2" if len(bullet.group(1)) >= 2 else "List Bullet"
            _inline(doc.add_paragraph(style=style), bullet.group(2))
            i += 1
            continue

        if numbered := _NUMBERED_RE.match(line):
            flush_paragraph()
            # The number is written literally: Word's auto-numbering keeps
            # one counter per style, so a second list would carry on
            # counting from the first.
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18 + 18 * (len(numbered.group(1)) // 2))
            _add_run(paragraph, f"{numbered.group(2)}. ")
            _inline(paragraph, numbered.group(3))
            i += 1
            continue

        paragraph_lines.append(line)
        i += 1

    flush_paragraph()
