"""Excel (``.xlsx``) workbook builder for exports.

The default spreadsheet export for a non-technical audience: opens with a
double-click, keeps multi-line text in one cell, and holds several
related tables (a coding's quotes, matrix and codebook) as sheets of one
file instead of several CSVs.

Two safety properties every sheet gets, whatever the caller passes:

* **No formulas.** openpyxl stores any string beginning with ``=`` as a
  live formula, so a quote reading ``=HYPERLINK(...)`` would become one.
  Every string cell is pinned to the string type instead (a formula-like
  one also gets Excel's ``quotePrefix`` text flag), and its text
  run through ``export_text.xml_safe`` (openpyxl raises on the control
  characters XML forbids, which a pasted vertical tab would otherwise
  turn into a 500).
* **No silent loss.** A cell over Excel's 32,767-character limit is cut
  short with a visible marker, and the count of cut cells is reported on
  the workbook's closing "About this export" sheet.

Output is byte-deterministic (see ``freeze_ooxml``) so a project bundle's
manifest hashes stay stable across identical exports.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from backend.app.core.export_text import neutralize_formula, xml_safe

# Matches the project bundle's fixed ZIP timestamp.
FIXED_TIMESTAMP = datetime(2026, 1, 1, tzinfo=timezone.utc)
_FIXED_ZIP_DATE = (2026, 1, 1, 0, 0, 0)
_FIXED_W3CDTF = "2026-01-01T00:00:00Z"
_CORE_TIMESTAMP_RE = re.compile(
    rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)

EXCEL_CELL_LIMIT = 32767
_TRUNCATION_MARKER = " … [truncated {n} characters]"
_MAX_COLUMN_WIDTH = 60
_WRAP_THRESHOLD = 40

ABOUT_SHEET_TITLE = "About this export"
CREATOR = "Qualitative Coding Tool"


def freeze_ooxml(data: bytes) -> bytes:
    """Rewrite an Office Open XML package so identical content always
    yields identical bytes.

    openpyxl stamps ``dcterms:modified`` with the wall clock on every
    save, and both openpyxl and python-docx write each ZIP entry with the
    current time. Entries keep their original order (``[Content_Types].xml``
    position included); only timestamps change.
    """
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, mode="w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = _CORE_TIMESTAMP_RE.sub(
                    lambda m: m.group(1) + _FIXED_W3CDTF.encode() + m.group(2), content
                )
            zinfo = zipfile.ZipInfo(filename=info.filename, date_time=_FIXED_ZIP_DATE)
            zinfo.compress_type = zipfile.ZIP_DEFLATED
            zinfo.external_attr = 0o644 << 16
            dst.writestr(zinfo, content)
    return out.getvalue()


@dataclass
class SheetSpec:
    """One worksheet: a fixed (never user-supplied) title, a header row,
    and data rows of ``str``/``int``/``float``/``None`` cells."""

    title: str
    headers: Sequence[str]
    rows: Sequence[Sequence[Any]]
    column_widths: dict[int, int] = field(default_factory=dict)


def _truncate(text: str) -> tuple[str, bool]:
    if len(text) <= EXCEL_CELL_LIMIT:
        return text, False
    # Sized against the widest possible count, so the marker always fits.
    keep = EXCEL_CELL_LIMIT - len(_TRUNCATION_MARKER.format(n=len(text)))
    return text[:keep] + _TRUNCATION_MARKER.format(n=len(text) - keep), True


def _write_cell(ws: Any, row: int, col: int, value: Any) -> tuple[int, bool]:
    """Write one cell; return its display length and whether it was cut."""
    cell = ws.cell(row=row, column=col)
    if value is None or value == "":
        return 0, False
    if isinstance(value, bool):
        value = int(value)
    if isinstance(value, (int, float)):
        cell.value = value
        return len(str(value)), False
    text, truncated = _truncate(xml_safe(value))
    cell.value = text
    # Pin the type *after* assignment: openpyxl's setter re-infers it and
    # would mark a leading "=" as a formula ("f").
    cell.data_type = "s"
    if neutralize_formula(text) != text:
        # Excel's own "this is text" flag: the cell shows without a
        # visible apostrophe and stays text even if a user edits it.
        cell.quotePrefix = True
    longest_line = max((len(line) for line in text.split("\n")), default=0)
    if "\n" in text or longest_line > _WRAP_THRESHOLD:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    else:
        cell.alignment = Alignment(vertical="top")
    return longest_line, truncated


def _fill_sheet(ws: Any, spec: SheetSpec) -> int:
    """Fill ``ws`` from ``spec``; return the number of truncated cells."""
    bold = Font(bold=True)
    widths = [0] * len(spec.headers)
    for col, header in enumerate(spec.headers, start=1):
        length, _ = _write_cell(ws, 1, col, header)
        ws.cell(row=1, column=col).font = bold
        widths[col - 1] = length

    truncated = 0
    for r, row in enumerate(spec.rows, start=2):
        for col, value in enumerate(row, start=1):
            length, cut = _write_cell(ws, r, col, value)
            truncated += cut
            if col <= len(widths):
                widths[col - 1] = max(widths[col - 1], length)

    for idx, width in enumerate(widths):
        override = spec.column_widths.get(idx)
        ws.column_dimensions[get_column_letter(idx + 1)].width = override or min(
            max(width + 2, 8), _MAX_COLUMN_WIDTH
        )
    if spec.headers:
        ws.freeze_panes = "A2"
        last = get_column_letter(len(spec.headers))
        ws.auto_filter.ref = f"A1:{last}{max(len(spec.rows) + 1, 1)}"
    return truncated


def build_workbook(sheets: Sequence[SheetSpec], about: Sequence[tuple[str, Any]]) -> bytes:
    """An ``.xlsx`` of ``sheets`` in order, closed by an "About this
    export" sheet listing ``about`` (label, value) pairs plus the count
    of cells cut at Excel's cell limit.
    """
    wb = Workbook()
    wb.remove(wb.active)
    truncated = 0
    for spec in sheets:
        truncated += _fill_sheet(wb.create_sheet(title=spec.title), spec)

    about_rows: list[list[Any]] = [[label, value] for label, value in about]
    about_rows.append([
        "Cells shortened",
        f"{truncated} cell(s) exceeded Excel's {EXCEL_CELL_LIMIT:,}-character limit and were cut short"
        if truncated
        else "None",
    ])
    _fill_sheet(
        wb.create_sheet(title=ABOUT_SHEET_TITLE),
        SheetSpec(title=ABOUT_SHEET_TITLE, headers=["Field", "Value"], rows=about_rows),
    )

    wb.properties.creator = CREATOR
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP
    buffer = io.BytesIO()
    wb.save(buffer)
    return freeze_ooxml(buffer.getvalue())
