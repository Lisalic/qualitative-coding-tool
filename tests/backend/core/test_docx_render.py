"""Unit tests for backend/app/core/docx_render.py."""

import io
import zipfile

from docx import Document

from backend.app.core import docx_render


def _render(markdown: str) -> Document:
    doc = docx_render.new_document("Title")
    docx_render.markdown_to_docx(doc, markdown)
    return Document(io.BytesIO(docx_render.to_bytes(doc)))


def _styled(doc) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text]


def test_block_structure_maps_to_word_styles() -> None:
    doc = _render(
        "# Findings\n\nFirst para\ncontinues.\n\n## Sub\n\n- one\n- two\n  - nested\n\n"
        "1. first\n2. second\n\n> quoted\n\n```\ncode line\n```\n"
    )
    assert _styled(doc) == [
        ("Title", "Title"),
        ("Heading 1", "Findings"),
        ("Normal", "First para continues."),
        ("Heading 2", "Sub"),
        ("List Bullet", "one"),
        ("List Bullet", "two"),
        ("List Bullet 2", "nested"),
        ("Normal", "1. first"),
        ("Normal", "2. second"),
        ("Quote", "quoted"),
        ("Normal", "code line"),
    ]


def test_inline_bold_italic_code_and_escapes() -> None:
    doc = _render("Some **bold**, *italic*, _also_, `code` and \\*not italic\\*.")
    paragraph = doc.paragraphs[-1]
    assert paragraph.text == "Some bold, italic, also, code and *not italic*."
    runs = {r.text: r for r in paragraph.runs}
    assert runs["bold"].bold
    assert runs["italic"].italic and runs["also"].italic
    assert runs["code"].font.name == "Courier New"
    assert not any(r.italic for r in paragraph.runs if "not italic" in r.text)


def test_pipe_tables_become_word_tables() -> None:
    doc = _render("| Code | Count |\n|---|---:|\n| A \\| B | **3** |\n| C |\n")
    table = doc.tables[0]
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Code", "Count"],
        ["A | B", "3"],
        ["C", ""],
    ]


def test_raw_html_stays_literal_text() -> None:
    doc = _render("<script>alert(1)</script> <img src=x onerror=alert(1)>")
    assert doc.paragraphs[-1].text == "<script>alert(1)</script> <img src=x onerror=alert(1)>"


def test_links_become_text_never_a_hyperlink() -> None:
    doc = docx_render.new_document("T")
    docx_render.markdown_to_docx(doc, "See [the site](javascript:alert(1)) now.")
    content = docx_render.to_bytes(doc)
    assert Document(io.BytesIO(content)).paragraphs[-1].text == "See the site (javascript:alert(1)) now."
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        assert b"hyperlink" not in zf.read("word/_rels/document.xml.rels").lower()
        assert b"<w:hyperlink" not in zf.read("word/document.xml")


def test_control_characters_do_not_raise() -> None:
    doc = docx_render.new_document("Title\x0b with tab")
    docx_render.add_paragraph(doc, "body\x00text", label="Label\x1f")
    docx_render.add_table(doc, ["H\x0b"], [["cell\x0c"]])
    docx_render.markdown_to_docx(doc, "# Head\x0bing")
    parsed = Document(io.BytesIO(docx_render.to_bytes(doc)))
    assert "Label: bodytext" in [p.text for p in parsed.paragraphs]
    assert parsed.tables[0].rows[1].cells[0].text == "cell"


def test_output_is_byte_identical_across_builds() -> None:
    def build() -> bytes:
        doc = docx_render.new_document("Same", ["sub"])
        docx_render.markdown_to_docx(doc, "# A\n\nB")
        return docx_render.to_bytes(doc)

    assert build() == build()
