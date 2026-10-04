"""Unit tests for backend/app/core/xlsx_render.py."""

import io
import zipfile

from openpyxl import load_workbook

from backend.app.core.xlsx_render import (
    ABOUT_SHEET_TITLE,
    EXCEL_CELL_LIMIT,
    SheetSpec,
    build_workbook,
)


def _build(rows: list[list], headers: list[str] | None = None) -> bytes:
    return build_workbook([SheetSpec("Data", headers or ["Value"], rows)], about=[("File", "f")])


def test_formula_text_is_stored_as_a_string_never_a_formula() -> None:
    payload = '=HYPERLINK("http://x","y")'
    content = _build([[payload], [" +1"], ["@SUM(A1)"]], headers=["=header"])

    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        sheet_xml = zf.read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet_xml

    ws = load_workbook(io.BytesIO(content))["Data"]
    cells = [c for row in ws.iter_rows() for c in row]
    assert [c.value for c in cells] == ["=header", payload, " +1", "@SUM(A1)"]
    assert all(c.data_type == "s" for c in cells)
    # Excel's own "text" flag, so editing the cell can't turn it into a formula.
    assert all(c.quotePrefix for c in cells)


def test_numbers_stay_numeric() -> None:
    ws = load_workbook(io.BytesIO(_build([[-5], [2.5], ["plain"]])))["Data"]
    values = [(c.value, c.data_type, c.quotePrefix) for c in ws["A"]][1:]
    assert values == [(-5, "n", False), (2.5, "n", False), ("plain", "s", False)]


def test_control_characters_are_stripped_not_raised() -> None:
    ws = load_workbook(io.BytesIO(_build([["pasted\x0bfrom\x00word"]])))["Data"]
    assert ws["A2"].value == "pastedfromword"


def test_over_long_cells_are_cut_with_a_marker_and_counted() -> None:
    workbook = load_workbook(io.BytesIO(_build([["a" * 40000], ["short"]])))
    cell = workbook["Data"]["A2"].value
    assert len(cell) <= EXCEL_CELL_LIMIT
    kept, marker = cell.split(" … [truncated ")
    assert len(kept) + int(marker.removesuffix(" characters]")) == 40000
    about = {row[0].value: row[1].value for row in workbook[ABOUT_SHEET_TITLE].iter_rows(min_row=2)}
    assert about["File"] == "f"
    assert about["Cells shortened"].startswith("1 cell(s)")


def test_sheets_keep_their_order_and_end_with_about() -> None:
    content = build_workbook(
        [SheetSpec("First", ["a"], []), SheetSpec("Second", ["b"], [])], about=[]
    )
    assert load_workbook(io.BytesIO(content)).sheetnames == ["First", "Second", ABOUT_SHEET_TITLE]


def test_output_is_byte_identical_across_builds() -> None:
    rows = [["x", 1], ["=y", 2]]
    first = build_workbook([SheetSpec("S", ["a", "b"], rows)], about=[("k", "v")])
    second = build_workbook([SheetSpec("S", ["a", "b"], rows)], about=[("k", "v")])
    assert first == second
    with zipfile.ZipFile(io.BytesIO(first)) as zf:
        assert {info.date_time for info in zf.infolist()} == {(2026, 1, 1, 0, 0, 0)}
        assert b"2026-01-01T00:00:00Z</dcterms:modified>" in zf.read("docProps/core.xml")
