"""Unit tests for backend/app/services/export_service.py.

Validates deterministic, RFC 4180 compliant CSV and JSON exports for codebooks,
coding entries, row memos, and frequency summaries. Covers owner scoping,
historical versions, empty artifacts, and format validity.
"""

import io
import json
import re
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker

from backend.app.core.exceptions import NotFoundError, ValidationAppError
from backend.app.database import File, User
from backend.app.services import export_service, version_service
from backend.app.storage_models import CodingEntry, RowMemo, Submission


DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docx(content: bytes):
    return Document(io.BytesIO(content))


def _docx_text(content: bytes) -> str:
    doc = _docx(content)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def _xlsx(content: bytes):
    return load_workbook(io.BytesIO(content))


def _sheet_rows(workbook, title: str) -> list[list]:
    return [[c.value for c in row] for row in workbook[title].iter_rows()]


@pytest.fixture()
def SessionLocal(async_sqlite_engine):
    return async_sessionmaker(async_sqlite_engine, expire_on_commit=False)


@pytest.fixture()
async def session(SessionLocal):
    async with SessionLocal() as s:
        yield s


@pytest.fixture()
async def user_id(session) -> int:
    user = User(email="export-service-test@example.com", password="hash")
    session.add(user)
    await session.commit()
    await session.refresh(user)
    return user.id


@pytest.fixture()
async def other_user_id(session) -> int:
    user = User(email="export-other-user@example.com", password="hash")
    session.add(user)
    await session.commit()
    await session.refresh(user)
    return user.id


async def _make_file(session, owner_id: int, filename: str, file_type: str) -> File:
    file_rec = File(
        user_id=owner_id,
        filename=filename,
        schemaname="test_schema",
        file_type=file_type,
    )
    session.add(file_rec)
    await session.commit()
    await session.refresh(file_rec)
    return file_rec


def _make_code_row(code_uid: str, name: str, position: int = 1, family_uid: str = "", family_name: str = "", **kwargs) -> dict:
    row = {
        "code_uid": code_uid,
        "name": name,
        "position": position,
        "family_uid": family_uid,
        "family_name": family_name,
        "body": kwargs.get("body", name),
        "definition": kwargs.get("definition", ""),
        "inclusion": kwargs.get("inclusion", ""),
        "exclusion": kwargs.get("exclusion", ""),
        "keywords": kwargs.get("keywords", ""),
        "example": kwargs.get("example", ""),
    }
    return row


@pytest.mark.asyncio
async def test_export_codebook_empty(session, user_id):
    file_rec = await _make_file(session, user_id, "my_cb.csv", "codebook")
    await version_service.commit_codebook_version(
        session, file_id=file_rec.id, author_user_id=user_id, origin="manual", codes=[]
    )

    csv_out, media, filename = await export_service.export_codebook(
        session, file_rec.id, user_id, export_format="csv"
    )
    assert media == "text/csv; charset=utf-8"
    assert "code_uid,name,family_uid,family_name" in csv_out
    assert filename == "my_cb_v1_codebook.csv"

    qdc_out, media_q, filename_q = await export_service.export_codebook(
        session, file_rec.id, user_id, export_format="qdc"
    )
    assert media_q == "application/xml; charset=utf-8"
    assert filename_q == "my_cb_v1_codebook.qdc"
    assert 'xmlns="urn:QDA-XML:codebook:1.0"' in qdc_out

    # Word is the default -- most users read a codebook rather than
    # import it, and it is what the dropdown leads with.
    default_out, default_media, default_name = await export_service.export_codebook(
        session, file_rec.id, user_id
    )
    assert default_media == DOCX_MEDIA
    assert default_name == "my_cb_v1_codebook.docx"
    assert "This codebook has no codes yet." in _docx_text(default_out)


@pytest.mark.asyncio
async def test_export_codebook_with_codes(session, user_id):
    file_rec = await _make_file(session, user_id, "codes.csv", "codebook")
    codes = [
        _make_code_row("c2", "Beta", position=2, family_uid="f1", family_name="Fam"),
        _make_code_row("c1", "Alpha", position=1, family_uid="f1", family_name="Fam"),
    ]
    await version_service.commit_codebook_version(
        session, file_id=file_rec.id, author_user_id=user_id, origin="manual", codes=codes
    )

    csv_out, _, _ = await export_service.export_codebook(session, file_rec.id, user_id, export_format="csv")
    lines = csv_out.strip().split("\r\n")
    assert len(lines) == 3  # header + 2 codes
    # Ordered by position (c1 then c2)
    assert lines[1].startswith("c1,Alpha")
    assert lines[2].startswith("c2,Beta")

    qdc_out, _, filename_q = await export_service.export_codebook(session, file_rec.id, user_id, export_format="qdc")
    assert filename_q == "codes_v1_codebook.qdc"

    # One non-codable family element holding both codes in position order.
    import xml.etree.ElementTree as ET

    ns = {"q": "urn:QDA-XML:codebook:1.0"}
    root = ET.fromstring(qdc_out)
    families = list(root.find("q:Codes", ns))
    assert len(families) == 1
    assert families[0].get("name") == "Fam"
    assert families[0].get("isCodable") == "false"
    assert [c.get("name") for c in families[0].findall("q:Code", ns)] == ["Alpha", "Beta"]


@pytest.mark.asyncio
async def test_export_codebook_historical_version(session, user_id):
    file_rec = await _make_file(session, user_id, "hist_cb.csv", "codebook")
    v1_codes = [_make_code_row("c1", "V1 Code", position=1)]
    await version_service.commit_codebook_version(
        session, file_id=file_rec.id, author_user_id=user_id, origin="manual", codes=v1_codes
    )
    v2_codes = [
        _make_code_row("c1", "V1 Code", position=1),
        _make_code_row("c2", "V2 Code", position=2),
    ]
    await version_service.commit_codebook_version(
        session, file_id=file_rec.id, author_user_id=user_id, origin="manual", codes=v2_codes
    )

    csv_v1, _, filename_v1 = await export_service.export_codebook(
        session, file_rec.id, user_id, version_no=1, export_format="csv"
    )
    assert "V2 Code" not in csv_v1
    assert "V1 Code" in csv_v1
    assert filename_v1 == "hist_cb_v1_codebook.csv"

    csv_v2, _, filename_v2 = await export_service.export_codebook(
        session, file_rec.id, user_id, version_no=2, export_format="csv"
    )
    assert "V2 Code" in csv_v2
    assert filename_v2 == "hist_cb_v2_codebook.csv"


@pytest.mark.asyncio
async def test_export_codebook_unauthorized_and_type_check(session, user_id, other_user_id):
    cb = await _make_file(session, user_id, "my_cb.csv", "codebook")
    with pytest.raises(NotFoundError):
        await export_service.export_codebook(session, cb.id, other_user_id)

    non_cb = await _make_file(session, user_id, "data.csv", "dataset")
    with pytest.raises(ValidationAppError):
        await export_service.export_codebook(session, non_cb.id, user_id)


@pytest.mark.asyncio
async def test_export_coding_with_entries(session, user_id):
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )

    entry = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="post_100",
        code="Theme Alpha",
        code_uid="uid_alpha",
        quote="A memorable quote",
        start_offset=10,
        end_offset=27,
        notes="Important note",
        coder="human",
        valid_from=1,
        valid_to=None,
    )
    session.add(entry)
    await session.commit()

    csv_out, media, filename = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="csv"
    )
    assert media == "text/csv; charset=utf-8"
    assert "file_id,version_no,entry_id,row_type,post_id,code_uid,code" in csv_out
    assert "A memorable quote" in csv_out
    assert "post_100" in csv_out
    assert filename == "coded_data_v1_segments_long.csv"

    json_out, media_j, filename_j = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="json"
    )
    assert media_j == "application/json; charset=utf-8"
    data = json.loads(json_out)
    assert data["file_id"] == coding_file.id
    assert data["layout"] == "long"
    assert len(data["entries"]) == 1
    assert data["entries"][0]["quote"] == "A memorable quote"
    assert data["entries"][0]["notes"] == "Important note"
    assert filename_j == "coded_data_v1_segments_long.json"


@pytest.mark.asyncio
async def test_export_memos(session, user_id):
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")
    memo = RowMemo(
        file_id=coding_file.id,
        row_type="submission",
        row_id="post_100",
        body="Analytic memo for post 100",
        author_user_id=user_id,
    )
    session.add(memo)
    await session.commit()

    csv_out, media, filename = await export_service.export_memos(
        session, coding_file.id, user_id, export_format="csv"
    )
    assert "memo_id,row_type,row_id,body" in csv_out
    assert "Analytic memo for post 100" in csv_out
    assert filename == "coded_data_memos.csv"

    md_out, media_m, filename_m = await export_service.export_memos(
        session, coding_file.id, user_id, export_format="md"
    )
    assert media_m == "text/markdown; charset=utf-8"
    assert filename_m == "coded_data_memos.md"
    assert "# coded_data.csv -- memos" in md_out
    assert "## submission post_100" in md_out
    assert "Analytic memo for post 100" in md_out


@pytest.mark.asyncio
async def test_export_memos_word_is_the_default_and_keeps_paragraphs(session, user_id):
    """A memo body is prose -- neither the Word default nor the markdown
    export may flatten its paragraph breaks the way a single CSV cell does.
    """
    coding_file = await _make_file(session, user_id, "prose.csv", "coding")
    session.add(
        RowMemo(
            file_id=coding_file.id,
            row_type="submission",
            row_id="p1",
            body="First paragraph.\n\nSecond paragraph.",
            author_user_id=user_id,
        )
    )
    await session.commit()

    default_out, default_media, default_name = await export_service.export_memos(
        session, coding_file.id, user_id
    )
    assert default_media == DOCX_MEDIA
    assert default_name == "prose_memos.docx"
    paragraphs = [p.text for p in _docx(default_out).paragraphs]
    assert "Post p1" in paragraphs
    assert paragraphs[-2:] == ["First paragraph.", "Second paragraph."]

    md_out, md_media, _ = await export_service.export_memos(
        session, coding_file.id, user_id, export_format="md"
    )
    assert md_media == "text/markdown; charset=utf-8"
    assert "First paragraph.\n\nSecond paragraph." in md_out


@pytest.mark.asyncio
async def test_export_memos_markdown_when_empty(session, user_id):
    raw_file = await _make_file(session, user_id, "no_memos.csv", "raw_data")
    md_out, _, _ = await export_service.export_memos(session, raw_file.id, user_id, export_format="md")
    assert "_No memos._" in md_out


@pytest.mark.asyncio
async def test_export_summary(session, user_id):
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")
    e1 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p1",
        code="Theme B",
        code_uid="b",
        quote="q1",
        start_offset=0,
        end_offset=2,
    )
    e2 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p2",
        code="Theme A",
        code_uid="a",
        quote="q2",
        start_offset=0,
        end_offset=2,
    )
    e3 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p3",
        code="Theme A",
        code_uid="a",
        quote="q3",
        start_offset=0,
        end_offset=2,
    )
    session.add_all([e1, e2, e3])
    await session.commit()

    md_out, media, filename = await export_service.export_summary(
        session, coding_file.id, user_id, export_format="md"
    )
    assert media == "text/markdown; charset=utf-8"
    assert filename == "coded_data_summary.md"
    md_lines = md_out.strip().split("\n")
    assert md_lines[0] == "# coded_data.csv -- code frequency"
    assert md_lines[2] == "| Code | Family | Frequency | Documents | Code UID |"
    # Descending frequency, then code_uid -- the repository's ordering.
    assert md_lines[4] == "| Theme A |  | 2 | 2 | a |"
    assert md_lines[5] == "| Theme B |  | 1 | 1 | b |"


@pytest.mark.asyncio
async def test_exports_reject_formats_the_artifact_does_not_offer(session, user_id):
    """An unrecognised format must fail loudly rather than fall through to
    whichever branch happens to be last -- asking a codebook for "md" used
    to hand back CSV content under a .csv filename.
    """
    codebook_file = await _make_file(session, user_id, "cb.csv", "codebook")
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")

    for bad in ("md", "json", ""):
        with pytest.raises(ValidationAppError):
            await export_service.export_codebook(session, codebook_file.id, user_id, export_format=bad)

    for bad in ("md", "qdc", ""):
        with pytest.raises(ValidationAppError):
            await export_service.export_coding(session, coding_file.id, user_id, export_format=bad)

    for bad in ("csv", "json", ""):
        with pytest.raises(ValidationAppError):
            await export_service.export_summary(session, coding_file.id, user_id, export_format=bad)

    for bad in ("json", "qdc", ""):
        with pytest.raises(ValidationAppError):
            await export_service.export_memos(session, coding_file.id, user_id, export_format=bad)


@pytest.mark.asyncio
async def test_export_summary_markdown_when_no_codes_applied(session, user_id):
    coding_file = await _make_file(session, user_id, "empty.csv", "coding")
    md_out, _, _ = await export_service.export_summary(session, coding_file.id, user_id, export_format="md")
    assert "_No codes applied._" in md_out


@pytest.mark.asyncio
async def test_repeated_exports_are_byte_identical(session, user_id):
    """Exports of unchanged data must produce byte-identical output (QC-001)."""
    coding_file = await _make_file(session, user_id, "repeat_test.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    e1 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p1",
        code="Theme A",
        code_uid="a",
        quote="q1",
        start_offset=0,
        end_offset=2,
        valid_from=1,
        valid_to=None,
    )
    memo = RowMemo(
        file_id=coding_file.id,
        row_type="submission",
        row_id="p1",
        body="Note",
        author_user_id=user_id,
    )
    session.add_all([e1, memo])
    await session.commit()

    # Repeated coding JSON exports
    c1, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="json")
    c2, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="json")
    assert c1 == c2

    # Repeated summary markdown exports
    s1, _, _ = await export_service.export_summary(session, coding_file.id, user_id, export_format="md")
    s2, _, _ = await export_service.export_summary(session, coding_file.id, user_id, export_format="md")
    assert s1 == s2

    # Repeated memo markdown exports
    m1, _, _ = await export_service.export_memos(session, coding_file.id, user_id, export_format="md")
    m2, _, _ = await export_service.export_memos(session, coding_file.id, user_id, export_format="md")
    assert m1 == m2


@pytest.mark.asyncio
async def test_export_summary_with_historical_version(session, user_id):
    """Historical summary exports respect entries_as_of version_no (QC-001)."""
    coding_file = await _make_file(session, user_id, "hist_summary.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    # Version 1 entry: Theme A (valid_from=1, valid_to=1)
    e_v1 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p1",
        code="Theme A",
        code_uid="a",
        quote="q1",
        start_offset=0,
        end_offset=2,
        valid_from=1,
        valid_to=1,
    )
    # Version 2 entries: Theme B (valid_from=2, valid_to=None), Theme B (valid_from=2, valid_to=None)
    e_v2_1 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p1",
        code="Theme B",
        code_uid="b",
        quote="q1",
        start_offset=0,
        end_offset=2,
        valid_from=2,
        valid_to=None,
    )
    e_v2_2 = CodingEntry(
        file_id=coding_file.id,
        row_type="submission",
        post_id="p2",
        code="Theme B",
        code_uid="b",
        quote="q2",
        start_offset=0,
        end_offset=2,
        valid_from=2,
        valid_to=None,
    )
    session.add_all([e_v1, e_v2_1, e_v2_2])
    await session.commit()
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )

    # Historical summary for version 1
    md_v1, _, fn_v1 = await export_service.export_summary(
        session, coding_file.id, user_id, version_no=1, export_format="md"
    )
    assert md_v1.splitlines()[0] == "# hist_summary.csv -- code frequency (v1)"
    assert "| Theme A |  | 1 | 1 | a |" in md_v1
    assert "Theme B" not in md_v1
    assert fn_v1 == "hist_summary_v1_summary.md"

    # Historical summary for version 2
    md_v2, _, fn_v2 = await export_service.export_summary(
        session, coding_file.id, user_id, version_no=2, export_format="md"
    )
    assert md_v2.splitlines()[0] == "# hist_summary.csv -- code frequency (v2)"
    assert "| Theme B |  | 2 | 2 | b |" in md_v2
    assert "Theme A" not in md_v2
    assert fn_v2 == "hist_summary_v2_summary.md"


@pytest.mark.asyncio
async def test_export_codebook_unknown_version_no_is_404(session, user_id):
    """A version_no that doesn't exist must 404, not silently export empty
    content as if the request had succeeded."""
    file_rec = await _make_file(session, user_id, "cb.csv", "codebook")
    await version_service.commit_codebook_version(
        session, file_id=file_rec.id, author_user_id=user_id, origin="manual", codes=[]
    )
    with pytest.raises(NotFoundError):
        await export_service.export_codebook(session, file_rec.id, user_id, version_no=99, export_format="csv")


@pytest.mark.asyncio
async def test_export_coding_unknown_version_no_is_404(session, user_id):
    coding_file = await _make_file(session, user_id, "coding.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    with pytest.raises(NotFoundError):
        await export_service.export_coding(session, coding_file.id, user_id, version_no=99, export_format="csv")


@pytest.mark.asyncio
async def test_export_summary_unknown_version_no_is_404(session, user_id):
    coding_file = await _make_file(session, user_id, "coding.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    with pytest.raises(NotFoundError):
        await export_service.export_summary(session, coding_file.id, user_id, version_no=99)


@pytest.mark.asyncio
async def test_export_coding_privacy_flags_default_off(session, user_id):
    """source_text and author are opt-in, not opt-out."""
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    session.add(Submission(file_id=coding_file.id, id="p1", title="t", selftext="secret body", author="alice", word_count=2))
    session.add(
        CodingEntry(
            file_id=coding_file.id, row_type="submission", post_id="p1", code="A", code_uid="a",
            quote="q", start_offset=0, end_offset=1, valid_from=1, valid_to=None,
        )
    )
    await session.commit()

    csv_out, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="csv")
    assert "source_text" not in csv_out
    assert "secret body" not in csv_out
    assert "author" not in csv_out
    assert "alice" not in csv_out

    csv_opt_in, _, _ = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="csv", include_source_text=True, include_author=True
    )
    assert "secret body" in csv_opt_in
    assert "alice" in csv_opt_in


@pytest.mark.asyncio
async def test_export_coding_wide_layout_includes_uncoded_rows(session, user_id):
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")
    codes = [
        _make_code_row("a", "Theme A", position=1),
        _make_code_row("b", "Theme B", position=2),
    ]
    await version_service.commit_codebook_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual", codes=codes
    )
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    # p1 is coded with "a"; p2 has no coding_entries at all -- wide format
    # must still surface it as an uncoded row (long format would omit it).
    session.add(Submission(file_id=coding_file.id, id="p1", title="t1", selftext="b1", word_count=1))
    session.add(Submission(file_id=coding_file.id, id="p2", title="t2", selftext="b2", word_count=1))
    session.add(
        CodingEntry(
            file_id=coding_file.id, row_type="submission", post_id="p1", code="Theme A", code_uid="a",
            quote="q", start_offset=0, end_offset=1, valid_from=1, valid_to=None,
        )
    )
    await session.commit()

    csv_out, _, filename = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="csv", layout="wide"
    )
    assert filename.endswith("_matrix_wide.csv")
    lines = csv_out.strip().split("\r\n")
    assert lines[0] == "\ufeffrow_type,post_id,is_coded,total_codes,code_a,code_b"
    rows_by_post = {line.split(",")[1]: line for line in lines[1:]}
    assert rows_by_post["p1"] == "submission,p1,1,1,1,0"
    assert rows_by_post["p2"] == "submission,p2,0,0,0,0"

    json_out, _, _ = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="json", layout="wide"
    )
    data = json.loads(json_out)
    assert data["layout"] == "wide"
    assert len(data["rows"]) == 2


@pytest.mark.asyncio
async def test_export_project_bundle_is_deterministic_and_privacy_scoped(session, user_id):
    from backend.app.database import Project, async_link_file_to_project

    project = Project(user_id=user_id, projectname="My Project", description="desc")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    cb_file = await _make_file(session, user_id, "cb.csv", "codebook")
    await version_service.commit_codebook_version(
        session, file_id=cb_file.id, author_user_id=user_id, origin="manual",
        codes=[_make_code_row("a", "Theme A", position=1)],
    )
    coding_file = await _make_file(session, user_id, "coding.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    session.add(Submission(file_id=coding_file.id, id="p1", title="t", selftext="body", author="alice", word_count=1))
    session.add(
        CodingEntry(
            file_id=coding_file.id, row_type="submission", post_id="p1", code="Theme A", code_uid="a",
            quote="q", start_offset=0, end_offset=1, valid_from=1, valid_to=None,
        )
    )
    await session.commit()

    await async_link_file_to_project(session, cb_file.id, project.id)
    await async_link_file_to_project(session, coding_file.id, project.id)
    await session.commit()

    bundle1, media, filename = await export_service.export_project_bundle(session, project.id, user_id)
    assert media == "application/zip"
    assert filename == "my_project_project_bundle.zip"

    bundle2, _, _ = await export_service.export_project_bundle(session, project.id, user_id)
    assert bundle1 == bundle2, "identical project state must produce byte-identical bundles"

    import zipfile
    import io as _io

    with zipfile.ZipFile(_io.BytesIO(bundle1)) as zf:
        names = zf.namelist()
        assert "manifest.json" in names
        assert "lineage/project_lineage.json" in names
        # Word/Excel by default, each artifact once.
        assert any(n.endswith("_codebook.docx") for n in names)
        assert not any(n.endswith(("_codebook.qdc", "_codebook.csv")) for n in names)
        assert any(n.endswith("_coding.xlsx") for n in names)
        assert not any(n.endswith((".csv", "_segments_long.json")) for n in names)

        manifest = json.loads(zf.read("manifest.json"))
        assert manifest["privacy_flags"] == {"include_source_text": False, "include_author": False}
        for entry in manifest["files"]:
            content = zf.read(entry["path"])
            assert len(content) == entry["bytes"]

        workbook = _xlsx(zf.read(next(n for n in names if n.endswith("_coding.xlsx"))))
        cells = {str(c.value) for ws in workbook.worksheets for row in ws.iter_rows() for c in row}
        assert "alice" not in cells
        assert "Author" not in cells and "Source text" not in cells


@pytest.mark.asyncio
async def test_export_project_bundle_writes_one_file_per_project_file(session, user_id):
    """One file in, one file out: no artifact may appear twice in two
    formats, and a file's memos ride along as a single ``.docx`` sidecar.
    """
    from backend.app.database import Project, async_link_file_to_project

    project = Project(user_id=user_id, projectname="One Each", description="")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    cb_file = await _make_file(session, user_id, "cb.csv", "codebook")
    await version_service.commit_codebook_version(
        session, file_id=cb_file.id, author_user_id=user_id, origin="manual",
        codes=[_make_code_row("a", "Theme A", position=1)],
    )
    coding_file = await _make_file(session, user_id, "coding.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    raw_file = await _make_file(session, user_id, "raw.csv", "raw_data")
    session.add_all(
        [
            RowMemo(file_id=coding_file.id, row_type="submission", row_id="p1", body="On the coding"),
            RowMemo(file_id=raw_file.id, row_type="submission", row_id="p1", body="On the raw data"),
        ]
    )
    await session.commit()

    for f in (cb_file, coding_file, raw_file):
        await async_link_file_to_project(session, f.id, project.id)
    await session.commit()

    bundle, _, _ = await export_service.export_project_bundle(session, project.id, user_id)

    import zipfile
    import io as _io

    with zipfile.ZipFile(_io.BytesIO(bundle)) as zf:
        names = zf.namelist()

    # No content is emitted in two formats.
    assert not any(n.endswith(("_memos.md", "_memos.csv", "_memos.xlsx")) for n in names)
    # `_slugify` drops the dot, so "coding.csv" slugs to "codingcsv".
    assert sorted(n for n in names if n.startswith("memos/")) == [
        f"memos/{coding_file.id}_codingcsv_memos.docx",
        f"memos/{raw_file.id}_rawcsv_memos.docx",
    ]

    # Each artifact appears exactly once, in its default format.
    assert [n for n in names if n.startswith("codebooks/")] == [f"codebooks/{cb_file.id}_cbcsv_codebook.docx"]
    assert [n for n in names if n.startswith("codings/")] == [f"codings/{coding_file.id}_codingcsv_coding.xlsx"]

    # Every artifact path is unique once its extension is stripped, so no
    # file is duplicated across formats anywhere in the bundle.
    artifact_paths = [n for n in names if n not in ("manifest.json", "lineage/project_lineage.json")]
    stems = [n.rsplit(".", 1)[0] for n in artifact_paths]
    assert len(stems) == len(set(stems))


@pytest.mark.asyncio
async def test_export_project_bundle_includes_comparison_markdown(session, user_id):
    from backend.app.database import Project, async_link_file_to_project

    project = Project(user_id=user_id, projectname="Cmp Project", description="")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    cmp_file = await _make_file(session, user_id, "A vs B", "codebook_comparison")
    await version_service.commit_blob_version(
        session, file_id=cmp_file.id, author_user_id=user_id, origin="generated",
        content="# Comparison\n\nA and B differ on...",
    )
    await async_link_file_to_project(session, cmp_file.id, project.id)
    await session.commit()

    bundle, _, _ = await export_service.export_project_bundle(
        session, project.id, user_id, comparison_formats=["md", "docx"]
    )

    import zipfile
    import io as _io

    with zipfile.ZipFile(_io.BytesIO(bundle)) as zf:
        names = zf.namelist()
        docx_paths = [n for n in names if n.startswith("comparisons/") and n.endswith("_comparison.docx")]
        assert len(docx_paths) == 1
        doc = _docx(zf.read(docx_paths[0]))
        assert [(p.style.name, p.text) for p in doc.paragraphs if p.text] == [
            ("Title", "A vs B"),
            ("Normal", "Comparison · Version 1"),
            ("Heading 1", "Comparison"),
            ("Normal", "A and B differ on..."),
        ]
        comparison_paths = [n for n in names if n.startswith("comparisons/") and n.endswith("_comparison.md")]
        assert len(comparison_paths) == 1
        assert zf.read(comparison_paths[0]).decode("utf-8") == "# Comparison\n\nA and B differ on..."

        manifest = json.loads(zf.read("manifest.json"))
        entry = next(e for e in manifest["files"] if e["path"] == comparison_paths[0])
        assert entry["media_type"] == "text/markdown"


@pytest.mark.asyncio
async def test_export_project_bundle_includes_saved_summary_markdown(session, user_id):
    from backend.app.database import Project, async_link_file_to_project

    project = Project(user_id=user_id, projectname="Sum Project", description="")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    sum_file = await _make_file(session, user_id, "Themes", "summary")
    await version_service.commit_blob_version(
        session, file_id=sum_file.id, author_user_id=user_id, origin="generated",
        content="# Themes\n\nThree themes.",
    )
    await async_link_file_to_project(session, sum_file.id, project.id)
    await session.commit()

    bundle, _, _ = await export_service.export_project_bundle(session, project.id, user_id)

    import zipfile
    import io as _io

    with zipfile.ZipFile(_io.BytesIO(bundle)) as zf:
        paths = [n for n in zf.namelist() if n.startswith("summaries/")]
        assert paths == [f"summaries/{sum_file.id}_themes_summary.docx"]
        assert "Three themes." in _docx_text(zf.read(paths[0]))

    md_bundle, _, _ = await export_service.export_project_bundle(
        session, project.id, user_id, summary_formats=["md"]
    )
    with zipfile.ZipFile(_io.BytesIO(md_bundle)) as zf:
        paths = [n for n in zf.namelist() if n.startswith("summaries/")]
        assert paths == [f"summaries/{sum_file.id}_themes_summary.md"]
        assert zf.read(paths[0]).decode("utf-8") == "# Themes\n\nThree themes."


@pytest.mark.asyncio
async def test_export_project_bundle_unowned_project_raises(session, user_id, other_user_id):
    from backend.app.core.exceptions import ForbiddenError
    from backend.app.database import Project

    project = Project(user_id=other_user_id, projectname="Not Mine")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    with pytest.raises(ForbiddenError):
        await export_service.export_project_bundle(session, project.id, user_id)


@pytest.mark.asyncio
async def test_export_project_bundle_writes_each_requested_format(session, user_id):
    """Asking for several formats of one type writes the same artifact
    once per format; types left unset keep their single default."""
    from backend.app.database import Project, async_link_file_to_project

    project = Project(user_id=user_id, projectname="Many Formats", description="")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    cb_file = await _make_file(session, user_id, "cb.csv", "codebook")
    await version_service.commit_codebook_version(
        session, file_id=cb_file.id, author_user_id=user_id, origin="manual",
        codes=[_make_code_row("a", "Theme A", position=1)],
    )
    coding_file = await _make_file(session, user_id, "coding.csv", "coding")
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    session.add(RowMemo(file_id=coding_file.id, row_type="submission", row_id="p1", body="A note"))
    await session.commit()
    for f in (cb_file, coding_file):
        await async_link_file_to_project(session, f.id, project.id)
    await session.commit()

    bundle, _, _ = await export_service.export_project_bundle(
        session, project.id, user_id,
        codebook_formats=["csv", "qdc", "csv"],
        coding_formats=["json", "csv_wide"],
    )

    import zipfile
    import io as _io

    with zipfile.ZipFile(_io.BytesIO(bundle)) as zf:
        names = zf.namelist()
        manifest = json.loads(zf.read("manifest.json"))

    assert sorted(n for n in names if n.startswith("codebooks/")) == [
        f"codebooks/{cb_file.id}_cbcsv_codebook.csv",
        f"codebooks/{cb_file.id}_cbcsv_codebook.qdc",
    ]
    assert sorted(n for n in names if n.startswith("codings/")) == [
        f"codings/{coding_file.id}_codingcsv_matrix_wide.csv",
        f"codings/{coding_file.id}_codingcsv_segments_long.json",
    ]
    assert [n for n in names if n.startswith("memos/")] == [f"memos/{coding_file.id}_codingcsv_memos.docx"]
    assert manifest["formats"] == {
        "codebook": ["qdc", "csv"],
        "coding": ["csv_wide", "json"],
        "comparison": ["docx"],
        "summary": ["docx"],
        "memos": ["docx"],
    }


@pytest.mark.asyncio
async def test_export_project_bundle_rejects_an_unknown_format(session, user_id):
    from backend.app.database import Project

    project = Project(user_id=user_id, projectname="Bad Format", description="")
    session.add(project)
    await session.commit()
    await session.refresh(project)

    with pytest.raises(ValidationAppError):
        await export_service.export_project_bundle(
            session, project.id, user_id, codebook_formats=["json"]
        )


async def _make_coded_file(session, user_id: int, quote: str = "q1") -> File:
    """A coding with two codes in a family, two posts (p2 uncoded), and an
    AI-coded entry on p1 carrying ``quote``."""
    coding_file = await _make_file(session, user_id, "coded_data.csv", "coding")
    codes = [
        _make_code_row("a", "Theme A", position=1, family_uid="f1", family_name="Family One",
                       definition="Talks about A"),
        _make_code_row("b", "Theme B", position=2, family_uid="f1", family_name="Family One"),
    ]
    await version_service.commit_codebook_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual", codes=codes
    )
    await version_service.commit_coding_version(
        session, file_id=coding_file.id, author_user_id=user_id, origin="manual"
    )
    session.add(Submission(file_id=coding_file.id, id="p1", title="t1", selftext="secret body",
                           author="alice", word_count=2))
    session.add(Submission(file_id=coding_file.id, id="p2", title="t2", selftext="b2", author="bob",
                           word_count=1))
    session.add(
        CodingEntry(
            file_id=coding_file.id, row_type="submission", post_id="p1", code="Theme A", code_uid="a",
            quote=quote, start_offset=0, end_offset=2, valid_from=1, valid_to=None,
            coder="ai", coder_model="test/model",
        )
    )
    await session.commit()
    return coding_file


@pytest.mark.asyncio
async def test_export_coding_xlsx_is_the_default_with_every_view_as_a_sheet(session, user_id):
    coding_file = await _make_coded_file(session, user_id)

    content, media, filename = await export_service.export_coding(session, coding_file.id, user_id)
    assert media == XLSX_MEDIA
    assert filename == "coded_data_v2_coding.xlsx"

    workbook = _xlsx(content)
    assert workbook.sheetnames == ["Coded quotes", "Matrix", "Codebook", "Code counts", "About this export"]

    quotes = _sheet_rows(workbook, "Coded quotes")
    assert quotes[0][:8] == ["Item type", "Item ID", "Code", "Family", "Quote", "Notes", "Coded by", "Model"]
    assert quotes[1][:8] == ["Post", "p1", "Theme A", "Family One", "q1", None, "AI", "test/model"]

    # The matrix keeps the uncoded item and heads code columns by name.
    matrix = _sheet_rows(workbook, "Matrix")
    assert matrix[0] == ["Item type", "Item ID", "Coded", "Number of codes", "Theme A", "Theme B"]
    assert matrix[1:] == [["Post", "p1", "Yes", 1, 1, 0], ["Post", "p2", "No", 0, 0, 0]]

    assert _sheet_rows(workbook, "Code counts")[1] == ["Theme A", "Family One", 1, 1, "a"]
    assert _sheet_rows(workbook, "Codebook")[1][:3] == ["Family One", "Theme A", "Talks about A"]


@pytest.mark.asyncio
async def test_export_coding_word_and_excel_keep_privacy_flags_off_by_default(session, user_id):
    coding_file = await _make_coded_file(session, user_id)

    xlsx, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="xlsx")
    workbook = _xlsx(xlsx)
    cells = {c.value for ws in workbook.worksheets for row in ws.iter_rows() for c in row}
    assert not cells & {"Author", "Source text", "alice", "bob"}
    assert not any("secret body" in str(v) for v in cells)

    docx, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="docx")
    text = _docx_text(docx)
    assert "alice" not in text and "secret body" not in text

    xlsx_on, _, _ = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="xlsx", include_author=True, include_source_text=True
    )
    quotes = _sheet_rows(_xlsx(xlsx_on), "Coded quotes")
    assert "Author" in quotes[0] and "Source text" in quotes[0]
    assert "alice" in quotes[1] and "t1\n\nsecret body" in quotes[1]

    docx_on, _, _ = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="docx", include_author=True
    )
    assert "Author: alice" in _docx_text(docx_on)


@pytest.mark.asyncio
async def test_export_coding_word_report_groups_quotes_under_codes(session, user_id):
    coding_file = await _make_coded_file(session, user_id)
    content, media, filename = await export_service.export_coding(
        session, coding_file.id, user_id, export_format="docx"
    )
    assert media == DOCX_MEDIA
    assert filename == "coded_data_v2_coded_quotes.docx"
    styled = [(p.style.name, p.text) for p in _docx(content).paragraphs if p.text]
    assert styled[styled.index(("Heading 1", "Family One")):] == [
        ("Heading 1", "Family One"),
        ("Heading 2", "Theme A"),
        ("Normal", "Definition: Talks about A"),
        ("Quote", "q1"),
        ("Normal", "Post p1 · coded by AI (test/model)"),
        ("Heading 2", "Theme B"),
        ("Normal", "No quotes coded."),
    ]


@pytest.mark.asyncio
async def test_export_coding_layout_only_applies_to_csv_and_json(session, user_id):
    coding_file = await _make_file(session, user_id, "c.csv", "coding")
    for fmt in ("xlsx", "docx"):
        with pytest.raises(ValidationAppError):
            await export_service.export_coding(session, coding_file.id, user_id, export_format=fmt, layout="long")
    with pytest.raises(ValidationAppError):
        await export_service.export_coding(session, coding_file.id, user_id, export_format="csv", layout="tall")


@pytest.mark.asyncio
async def test_export_codebook_word_and_excel(session, user_id):
    file_rec = await _make_file(session, user_id, "cb.csv", "codebook")
    await version_service.commit_codebook_version(
        session, file_id=file_rec.id, author_user_id=user_id, origin="manual",
        codes=[
            _make_code_row("a", "Theme A", position=1, family_uid="f1", family_name="Fam",
                           definition="Def A", inclusion="Inc A"),
            _make_code_row("b", "Loose", position=2),
        ],
    )

    docx, _, _ = await export_service.export_codebook(session, file_rec.id, user_id)
    styled = [(p.style.name, p.text) for p in _docx(docx).paragraphs if p.text]
    assert styled[3:] == [
        ("Heading 1", "Fam"),
        ("Heading 2", "Theme A"),
        ("Normal", "Definition: Def A"),
        ("Normal", "Include when: Inc A"),
        ("Heading 1", "Codes without a family"),
        ("Heading 2", "Loose"),
    ]

    xlsx, media, filename = await export_service.export_codebook(session, file_rec.id, user_id, export_format="xlsx")
    assert (media, filename) == (XLSX_MEDIA, "cb_v1_codebook.xlsx")
    rows = _sheet_rows(_xlsx(xlsx), "Codebook")
    assert rows[0] == ["Family", "Code", "Definition", "Include when", "Exclude when", "Example", "Keywords", "Code ID"]
    assert rows[1] == ["Fam", "Theme A", "Def A", "Inc A", None, None, None, "a"]


@pytest.mark.asyncio
async def test_exports_carry_injection_payloads_as_inert_text(session, user_id):
    """Quotes, memos and code names are user/Reddit text: a formula must
    stay text in every spreadsheet format, a pasted control character must
    not break any XML format, and a code name must not inject markdown."""
    payload = '=HYPERLINK("http://evil.example","click")'
    coding_file = await _make_coded_file(session, user_id, quote=payload + "\x0b")
    session.add(RowMemo(file_id=coding_file.id, row_type="submission", row_id="p1",
                        body="@SUM(1)\x0b memo", author_user_id=user_id))
    await session.commit()

    xlsx, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="xlsx")
    with zipfile.ZipFile(io.BytesIO(xlsx)) as zf:
        assert not any(b"<f>" in zf.read(n) for n in zf.namelist() if n.startswith("xl/worksheets/"))
    quote_cell = _xlsx(xlsx)["Coded quotes"]["E2"]
    assert (quote_cell.value, quote_cell.data_type) == (payload, "s")

    memo_xlsx, _, _ = await export_service.export_memos(session, coding_file.id, user_id, export_format="xlsx")
    memo_cell = _xlsx(memo_xlsx)["Memos"]["C2"]
    assert (memo_cell.value, memo_cell.data_type) == ("@SUM(1) memo", "s")

    csv_out, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="csv")
    assert "'=HYPERLINK" in csv_out

    docx, _, _ = await export_service.export_coding(session, coding_file.id, user_id, export_format="docx")
    assert payload in _docx_text(docx)
    memo_docx, _, _ = await export_service.export_memos(session, coding_file.id, user_id)
    assert "@SUM(1) memo" in _docx_text(memo_docx)


@pytest.mark.asyncio
async def test_markdown_headings_and_cells_escape_user_text(session, user_id):
    coding_file = await _make_file(session, user_id, "evil\n# <script>x</script>", "coding")
    session.add(CodingEntry(
        file_id=coding_file.id, row_type="submission", post_id="p1", code="<img src=x onerror=alert(1)>|x",
        code_uid="z", quote="q", start_offset=0, end_offset=1,
    ))
    await session.commit()

    md_out, _, _ = await export_service.export_summary(session, coding_file.id, user_id, export_format="md")
    lines = md_out.split("\n")
    # The newline is flattened, so the "#" after it is plain text mid-line.
    assert lines[0] == "# evil # \\<script\\>x\\</script\\> -- code frequency"
    assert "| \\<img src=x onerror=alert(1)\\>\\|x |  | 1 | 1 | z |" in lines
    assert re.search(r"(?<!\\)<", md_out) is None, "every < must be backslash-escaped"


class TestCsvSerialize:
    def test_neutralizes_formula_cells_but_not_numbers(self) -> None:
        out = export_service._csv_serialize(
            [["=HYPERLINK(\"http://x\")", "+1 agree", "-5 karma", "@mention", "plain", -5]],
            ["a", "b", "c", "d", "e", "score"],
        )
        row = out.split("\r\n")[1]
        assert row.startswith("\"'=HYPERLINK(\"\"http://x\"\")\",'+1 agree,'-5 karma,'@mention,plain,-5")

    def test_neutralizes_formulas_hidden_behind_leading_whitespace(self) -> None:
        out = export_service._csv_serialize([[" =1+1", "\n=1+1"]], ["a", "b"])
        assert out.split("\r\n", 1)[1].startswith("' =1+1,\"'\n=1+1\"")

    def test_starts_with_a_utf8_bom(self) -> None:
        assert export_service._csv_serialize([], ["a"]).startswith("﻿a")
